import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from allapp.accounts.models import AuditEvent, UserRoleScope
from allapp.baseinfo.models import Owner, OwnerWarehouseBinding
from allapp.locations.models import Warehouse

from .excel_import import IMPORT_SHEET_NAME, PACKAGE_SHEET_NAME, PRODUCT_HEADERS
from .models import Product, ProductCategory, ProductPackage, ProductUom


class AllProductExportApiTests(TestCase):
    url = "/api/products/export-all-excel/"

    @classmethod
    def setUpTestData(cls):
        cls.owner = Owner.objects.create(code="EXPORT-A", name="货主甲")
        cls.other_owner = Owner.objects.create(code="EXPORT-B", name="货主乙")
        cls.uom = ProductUom.objects.create(code="ALL-EA", name="导出基本单位")
        cls.carton_uom = ProductUom.objects.create(code="ALL-CTN", name="导出包装单位")
        cls.category = ProductCategory.objects.create(code="ALL-CATEGORY", name="导出分类")
        product_ct = ContentType.objects.get_for_model(Product)
        cls.view_permission = Permission.objects.get(
            content_type=product_ct, codename="view_product"
        )
        User = get_user_model()
        cls.user = User.objects.create_user(username="all-export-owner")
        cls.user.user_permissions.add(cls.view_permission)
        UserRoleScope.objects.create(
            user=cls.user, role=UserRoleScope.Role.OWNER_MANAGER, owner=cls.owner
        )
        cls.global_user = User.objects.create_user(username="all-export-global")
        cls.global_user.user_permissions.add(
            cls.view_permission,
            Permission.objects.get(content_type=product_ct, codename="view_all_owner_products"),
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def product(self, code, *, owner=None, **fields):
        return Product.objects.create(
            owner=owner or self.owner,
            code=code,
            name=fields.pop("name", code),
            category=self.category,
            base_uom=self.uom,
            **fields,
        )

    def workbook(self, *, user=None):
        if user is not None:
            self.client.force_authenticate(user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content), data_only=False)

    @staticmethod
    def rows(workbook, sheet_name=IMPORT_SHEET_NAME):
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        return [
            dict(zip(headers, row))
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        ]

    def test_owner_export_contains_all_visible_archives_and_packages(self):
        product = self.product(
            "0001",
            name="=商品名称",
            gtin="00012345",
            price=Decimal("19.90"),
            purchase_price=Decimal("11.25"),
            is_active=False,
        )
        self.product("0002", gtin=None)
        self.product("OTHER", owner=self.other_owner)
        self.product("DELETED", is_deleted=True)
        ProductPackage.objects.create(
            product=product, uom=self.uom, qty_in_base=1, barcode="000001", sort_order=2
        )
        carton_package = ProductPackage.objects.create(
            product=product,
            uom=self.carton_uom,
            qty_in_base=12,
            barcode="000012",
            is_purchase_default=True,
            sort_order=1,
        )
        carton_package.is_active = False
        carton_package.save(update_fields=["is_active"])

        response, workbook = self.workbook()

        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        self.assertEqual(
            tuple(cell.value for cell in workbook[IMPORT_SHEET_NAME][1]), PRODUCT_HEADERS
        )
        rows = self.rows(workbook)
        self.assertEqual([row["货主商品编码"] for row in rows], ["0001", "0002"])
        self.assertEqual(rows[0]["启用"], "否")
        self.assertEqual(rows[0]["标准贸易条码"], "00012345")
        self.assertEqual(Decimal(str(rows[0]["默认价格"])), Decimal("19.90"))
        self.assertEqual(Decimal(str(rows[0]["进价"])), Decimal("11.25"))
        name_column = PRODUCT_HEADERS.index("商品名称") + 1
        self.assertEqual(workbook[IMPORT_SHEET_NAME].cell(2, name_column).data_type, "s")
        packages = self.rows(workbook, PACKAGE_SHEET_NAME)
        self.assertEqual([row["包装条码"] for row in packages], ["000012", "000001"])
        self.assertEqual(packages[0]["包装换算数量"], 12)
        self.assertEqual(packages[0]["启用"], "否")
        self.assertEqual(packages[0]["采购默认"], "是")
        self.assertEqual(workbook["基础资料"]["A2"].value, self.owner.code)
        self.assertIsNone(workbook["基础资料"]["A3"].value)
        audit = AuditEvent.objects.get(action="products.export_all_excel", actor=self.user)
        self.assertTrue(audit.succeeded)
        self.assertEqual(audit.metadata["product_count"], 2)
        self.assertEqual(audit.metadata["package_count"], 2)

    def test_global_export_keeps_owner_identity_and_excludes_unavailable_owners(self):
        other_product = self.product("SAME", owner=self.other_owner)
        self.product("SAME")
        ProductPackage.objects.create(product=other_product, uom=self.carton_uom, qty_in_base=24)
        for index, fields in enumerate(({"is_active": False}, {"is_deleted": True})):
            owner = Owner.objects.create(
                code=f"EXCLUDED-{index}", name=f"无效货主{index}", **fields
            )
            self.product(f"EXCLUDED-{index}", owner=owner)

        _, workbook = self.workbook(user=self.global_user)

        self.assertEqual(
            [(row["货主编码"], row["货主商品编码"]) for row in self.rows(workbook)],
            [(self.owner.code, "SAME"), (self.other_owner.code, "SAME")],
        )
        self.assertEqual(
            self.rows(workbook, PACKAGE_SHEET_NAME)[0]["货主编码"], self.other_owner.code
        )
        self.assertEqual(workbook["基础资料"]["A2"].value, self.owner.code)
        self.assertEqual(workbook["基础资料"]["A3"].value, self.other_owner.code)
        self.assertIsNone(workbook["基础资料"]["A4"].value)

    def test_warehouse_export_is_limited_to_active_owner_bindings(self):
        warehouse = Warehouse.objects.create(code="ALL-WH", name="导出仓库")
        warehouse_user = get_user_model().objects.create_user(username="all-export-warehouse")
        warehouse_user.user_permissions.add(self.view_permission)
        UserRoleScope.objects.create(
            user=warehouse_user, role=UserRoleScope.Role.WAREHOUSE_MANAGER, warehouse=warehouse
        )
        OwnerWarehouseBinding.objects.create(owner=self.owner, warehouse=warehouse)
        OwnerWarehouseBinding.objects.create(
            owner=self.other_owner, warehouse=warehouse, is_active=False
        )
        self.product("VISIBLE")
        self.product("HIDDEN", owner=self.other_owner)

        _, workbook = self.workbook(user=warehouse_user)

        self.assertEqual([row["货主商品编码"] for row in self.rows(workbook)], ["VISIBLE"])

    def test_requires_authentication_view_permission_and_valid_scope(self):
        self.assertEqual(APIClient().get(self.url).status_code, 401)
        denied = get_user_model().objects.create_user(username="all-export-denied")
        UserRoleScope.objects.create(
            user=denied, role=UserRoleScope.Role.OWNER_MANAGER, owner=self.owner
        )
        unscoped = get_user_model().objects.create_user(username="all-export-unscoped")
        unscoped.user_permissions.add(self.view_permission)
        for user in (denied, unscoped):
            self.client.force_authenticate(user)
            self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertFalse(AuditEvent.objects.filter(action="products.export_all_excel").exists())

    def test_empty_export_is_valid_and_audited(self):
        _, workbook = self.workbook()

        self.assertEqual(self.rows(workbook), [])
        self.assertEqual(self.rows(workbook, PACKAGE_SHEET_NAME), [])
        audit = AuditEvent.objects.get(action="products.export_all_excel", actor=self.user)
        self.assertEqual(audit.metadata["product_count"], 0)
        self.assertEqual(audit.metadata["package_count"], 0)

    def test_export_is_not_limited_by_the_import_row_limit(self):
        for index in range(1001):
            self.product(f"BULK-{index:04d}", name=f"商品 {index}")

        _, workbook = self.workbook()

        self.assertEqual(len(self.rows(workbook)), 1001)
        self.assertEqual(self.rows(workbook)[-1]["货主商品编码"], "BULK-1000")
        audit = AuditEvent.objects.get(action="products.export_all_excel", actor=self.user)
        self.assertEqual(audit.metadata["product_count"], 1001)
