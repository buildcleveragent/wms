"""PDA return CRUD, catalog, statistics and exports under one authoritative scope."""

from decimal import Decimal
from io import BytesIO

from django.db.models import Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import serializers, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.renderers import JSONRenderer
from rest_framework.response import Response
from rest_framework.utils.encoders import JSONEncoder

from allapp.products.identifier_lookup import filter_by_product_search

from .assisted_history import _aware_bound, parse_period
from .return_models import ReturnEvent, ReturnOrder
from .return_services import (
    delete_draft,
    orders_for,
    post_return,
    return_scope,
    stock_payload,
    stock_queryset,
    write_draft,
)


class QuantityJSONEncoder(JSONEncoder):
    """Keep four-decimal inventory quantities exact across the JavaScript boundary."""

    def default(self, obj):
        if isinstance(obj, Decimal):
            return format(obj, "f")
        return super().default(obj)


class QuantityJSONRenderer(JSONRenderer):
    encoder_class = QuantityJSONEncoder


class LineInput(serializers.Serializer):
    inventory_id = serializers.IntegerField(min_value=1)
    available_qty = serializers.DecimalField(
        max_digits=18, decimal_places=4, min_value=Decimal(0), default=Decimal(0)
    )
    damaged_qty = serializers.DecimalField(
        max_digits=18, decimal_places=4, min_value=Decimal(0), default=Decimal(0)
    )
    new_damage_qty = serializers.DecimalField(
        max_digits=18, decimal_places=4, min_value=Decimal(0), default=Decimal(0)
    )


class DraftInput(serializers.Serializer):
    request_id = serializers.UUIDField(required=False)
    version = serializers.IntegerField(min_value=1, required=False)
    owner_id = serializers.IntegerField(min_value=1)
    recipient_type = serializers.ChoiceField(choices=ReturnOrder.Recipient.choices)
    supplier_id = serializers.IntegerField(
        min_value=1, required=False, allow_null=True, default=None
    )
    reason = serializers.CharField(max_length=200)
    lines = LineInput(many=True, allow_empty=False, max_length=200)

    def validate(self, data):
        if (data["recipient_type"] == "SUPPLIER") != bool(data["supplier_id"]):
            raise serializers.ValidationError("退供应商必须指定供应商，退货主不能指定供应商。")
        required = "version" if self.context.get("editing") else "request_id"
        if required not in data:
            raise serializers.ValidationError({required: "此字段必填。"})
        return data


class ActionInput(serializers.Serializer):
    request_id = serializers.UUIDField()
    version = serializers.IntegerField(min_value=1)
    reason = serializers.CharField(max_length=200, required=False)


class VersionInput(serializers.Serializer):
    version = serializers.IntegerField(min_value=1)


class FilterInput(serializers.Serializer):
    owner_id = serializers.IntegerField(min_value=1, required=False)
    supplier_id = serializers.IntegerField(min_value=1, required=False)
    product_id = serializers.IntegerField(min_value=1, required=False)
    operator_id = serializers.IntegerField(min_value=1, required=False)
    recipient_type = serializers.ChoiceField(choices=ReturnOrder.Recipient.choices, required=False)
    status = serializers.ChoiceField(choices=ReturnOrder.Status.choices, required=False)
    search = serializers.CharField(required=False, allow_blank=True, max_length=200)


def validated(serializer_class, data, **kwargs):
    serializer = serializer_class(data=data, **kwargs)
    serializer.is_valid(raise_exception=True)
    return serializer.validated_data


def order_payload(order, user):
    return {
        "id": order.pk,
        "return_no": order.return_no,
        "status": order.status,
        "version": order.version,
        "owner_id": order.owner_id,
        "owner_name": order.owner.name,
        "recipient_type": order.recipient_type,
        "supplier_id": order.supplier_id,
        "recipient_name": order.recipient_name,
        "operator_id": order.operator_id,
        "operator_name": order.operator.get_full_name() or order.operator.username,
        "reason": order.reason,
        "reversal_reason": order.reversal_reason,
        "created_at": order.created_at,
        "confirmed_at": order.confirmed_at,
        "reversed_at": order.reversed_at,
        "can_edit": order.operator_id == user.pk and order.status == "DRAFT",
        "can_reverse": order.operator_id == user.pk and order.status == "COMPLETED",
        "lines": [
            {
                "id": line.pk,
                "inventory_id": line.inventory_id,
                "product_id": line.product_id,
                "product_name": line.product_name,
                "product_code": line.product_code,
                "spec": line.spec,
                "base_unit": line.base_unit,
                "location_name": line.location_name,
                "dimensions": line.dimensions,
                "supplier_id": line.source_supplier_id,
                "supplier_name": line.source_supplier_name,
                "available_qty": line.available_qty,
                "damaged_qty": line.damaged_qty,
                "new_damage_qty": line.new_damage_qty,
            }
            for line in order.lines.all()
        ],
    }


def filtered_orders(user, params):
    data = validated(FilterInput, params)
    qs = orders_for(user)
    for name in ("owner_id", "supplier_id", "operator_id", "recipient_type", "status"):
        if name in data:
            qs = qs.filter(**{name: data[name]})
    if data.get("product_id"):
        qs = qs.filter(lines__product_id=data["product_id"])
    if data.get("search"):
        qs = qs.filter(
            Q(return_no__icontains=data["search"])
            | Q(owner__name__icontains=data["search"])
            | Q(recipient_name__icontains=data["search"])
        )
    return qs.distinct(), data


def event_queryset(user, params):
    orders, filters = filtered_orders(user, params)
    start, end = parse_period(params, default_today=True, max_days=366)
    qs = ReturnEvent.objects.filter(
        line__order__in=orders,
        kind__in=["CONFIRM", "REVERSE"],
        occurred_at__gte=_aware_bound(start),
        occurred_at__lt=_aware_bound(end, end=True),
    )
    if filters.get("product_id"):
        qs = qs.filter(line__product_id=filters["product_id"])
    return (
        qs.select_related("line__order__owner", "line__order__operator", "actor").order_by(
            "occurred_at", "id"
        ),
        start,
        end,
    )


def summarize(events):
    groups = {
        name: {} for name in ("units", "daily", "owners", "recipients", "products", "operators")
    }
    confirmed, reversed_ids = set(), set()
    for event in events:
        line, order = event.line, event.line.order
        is_reverse = event.kind == "REVERSE"
        (reversed_ids if is_reverse else confirmed).add(order.pk)
        at = (
            timezone.localtime(event.occurred_at)
            if timezone.is_aware(event.occurred_at)
            else event.occurred_at
        )
        keys = {
            "units": (line.base_unit, line.base_unit),
            "daily": (str(at.date()), str(at.date())),
            "owners": (order.owner_id, order.owner.name),
            "recipients": (
                f"{order.recipient_type}:{order.supplier_id or order.owner_id}",
                order.recipient_name,
            ),
            "products": (line.product_id, f"{line.product_name} · {order.owner.name}"),
            "operators": (
                order.operator_id,
                order.operator.get_full_name() or order.operator.username,
            ),
        }
        for category, (identity, label) in keys.items():
            key = (identity, line.base_unit)
            row = groups[category].setdefault(
                key,
                {
                    "id": identity,
                    "label": label,
                    "base_unit": line.base_unit,
                    "normal_qty": Decimal(0),
                    "damaged_qty": Decimal(0),
                    "new_damage_qty": Decimal(0),
                    "reverse_qty": Decimal(0),
                    "net_qty": Decimal(0),
                    "confirmed": set(),
                    "reversed": set(),
                },
            )
            total = line.available_qty + line.damaged_qty + line.new_damage_qty
            row["reversed" if is_reverse else "confirmed"].add(order.pk)
            if is_reverse:
                row["reverse_qty"] += total
                row["net_qty"] -= total
            else:
                row["normal_qty"] += line.available_qty
                row["damaged_qty"] += line.damaged_qty + line.new_damage_qty
                row["new_damage_qty"] += line.new_damage_qty
                row["net_qty"] += total
    result = {"confirmed_count": len(confirmed), "reversed_count": len(reversed_ids)}
    for name, rows in groups.items():
        result[name] = []
        for row in rows.values():
            row["confirmed_count"] = len(row.pop("confirmed"))
            row["reversed_count"] = len(row.pop("reversed"))
            result[name].append(row)
    return result


class Pagination(PageNumberPagination):
    page_size = 30
    page_size_query_param = "page_size"
    max_page_size = 100


class ReturnOrderViewSet(viewsets.GenericViewSet):
    serializer_class = DraftInput
    lookup_value_regex = r"\d+"
    permission_classes = [IsAuthenticated]
    renderer_classes = [QuantityJSONRenderer]
    pagination_class = Pagination

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        return_scope(request.user)

    def get_queryset(self):
        return (
            orders_for(self.request.user)
            .select_related("owner", "operator")
            .prefetch_related("lines")
        )

    def list(self, request):
        qs, _ = filtered_orders(request.user, request.query_params)
        start, end = parse_period(request.query_params, max_days=366)
        if start:
            qs = qs.filter(
                Q(
                    confirmed_at__gte=_aware_bound(start),
                    confirmed_at__lt=_aware_bound(end, end=True),
                )
                | Q(
                    reversed_at__gte=_aware_bound(start),
                    reversed_at__lt=_aware_bound(end, end=True),
                )
                | Q(
                    status="DRAFT",
                    created_at__gte=_aware_bound(start),
                    created_at__lt=_aware_bound(end, end=True),
                )
            )
        rows = self.paginate_queryset(
            qs.select_related("owner", "operator").prefetch_related("lines").order_by("-id")
        )
        return self.get_paginated_response([order_payload(row, request.user) for row in rows])

    def retrieve(self, request, pk=None):
        return Response(order_payload(self.get_object(), request.user))

    def create(self, request):
        data = validated(DraftInput, request.data)
        return Response(order_payload(write_draft(request.user, data), request.user), status=201)

    def partial_update(self, request, pk=None):
        self.get_object()
        data = validated(DraftInput, request.data, context={"editing": True})
        return Response(order_payload(write_draft(request.user, data, int(pk)), request.user))

    def destroy(self, request, pk=None):
        data = validated(VersionInput, request.data)
        delete_draft(request.user, int(pk), data["version"])
        return Response(status=204)

    @action(detail=True, methods=["post"])
    def confirm(self, request, pk=None):
        self.get_object()
        data = validated(ActionInput, request.data)
        return Response(order_payload(post_return(request.user, int(pk), data), request.user))

    @action(detail=True, methods=["post"])
    def reverse(self, request, pk=None):
        self.get_object()
        data = validated(ActionInput, request.data)
        if not data.get("reason"):
            raise ValidationError({"reason": "冲销必须填写原因。"})
        return Response(
            order_payload(post_return(request.user, int(pk), data, reverse=True), request.user)
        )

    @action(detail=False, methods=["get"])
    def products(self, request):
        filters = validated(FilterInput, request.query_params)
        qs = filter_by_product_search(stock_queryset(request.user), filters.get("search", ""))
        if filters.get("owner_id"):
            qs = qs.filter(owner_id=filters["owner_id"])
        qs = qs.filter(Q(available_qty__gt=0) | Q(damaged_qty__gt=0))
        rows = (
            qs.values(
                "product_id",
                "product__name",
                "product__code",
                "product__spec",
                "owner_id",
                "owner__name",
                "base_unit",
            )
            .annotate(available_qty=Sum("available_qty"), damaged_qty=Sum("damaged_qty"))
            .order_by("product__name", "owner_id", "product_id", "base_unit")
        )
        page = self.paginate_queryset(rows)
        return self.get_paginated_response(
            [
                {
                    "product_id": row["product_id"],
                    "product_name": row["product__name"],
                    "product_code": row["product__code"],
                    "spec": row["product__spec"],
                    "owner_id": row["owner_id"],
                    "owner_name": row["owner__name"],
                    "base_unit": row["base_unit"],
                    "available_qty": row["available_qty"],
                    "damaged_qty": row["damaged_qty"],
                }
                for row in page
            ]
        )

    @action(detail=False, methods=["get"], url_path="stock-options")
    def stock_options(self, request):
        filters = validated(FilterInput, request.query_params)
        if not filters.get("product_id"):
            raise ValidationError({"product_id": "请选择商品。"})
        qs = (
            stock_queryset(request.user)
            .filter(product_id=filters["product_id"])
            .filter(Q(available_qty__gt=0) | Q(damaged_qty__gt=0))
            .order_by("location_id", "expiry_date", "id")
        )
        return self.get_paginated_response(
            [stock_payload(row) for row in self.paginate_queryset(qs)]
        )

    @action(detail=False, methods=["get"])
    def options(self, request):
        qs = orders_for(request.user)
        return Response(
            {
                "owners": list(
                    stock_queryset(request.user)
                    .values("owner_id", "owner__name")
                    .order_by("owner_id")
                    .distinct()
                ),
                "suppliers": list(
                    qs.filter(supplier__isnull=False)
                    .values("supplier_id", "recipient_name")
                    .order_by("supplier_id", "recipient_name")
                    .distinct()
                ),
                "operators": list(
                    qs.values("operator_id", "operator__username")
                    .order_by("operator_id")
                    .distinct()
                ),
                "products": list(
                    qs.values("lines__product_id", "lines__product_name")
                    .order_by("lines__product_id", "lines__product_name")
                    .distinct()
                ),
            }
        )

    @action(detail=False, methods=["get"])
    def stats(self, request):
        events, start, end = event_queryset(request.user, request.query_params)
        return Response({"start_date": start, "end_date": end, **summarize(events.iterator())})

    @action(detail=False, methods=["get"])
    def export(self, request):
        events, start, end = event_queryset(request.user, request.query_params)
        if events.count() > 50000:
            raise ValidationError("导出超过50000行，请缩小日期或筛选范围。")
        events = list(events)
        wb = Workbook()
        sheet = wb.active
        sheet.title = "筛选条件"
        sheet.append(["开始", str(start), "结束", str(end)])
        for key, value in validated(FilterInput, request.query_params).items():
            sheet.append([key, safe_cell(str(value))])
        summary = summarize(events)
        for category, title in [
            ("units", "单位汇总"),
            ("daily", "每日汇总"),
            ("owners", "货主汇总"),
            ("recipients", "接收方汇总"),
            ("products", "商品汇总"),
            ("operators", "操作员汇总"),
        ]:
            sheet = wb.create_sheet(title)
            sheet.append(
                [
                    "名称",
                    "基本单位",
                    "退货单数",
                    "冲销单数",
                    "正常退货",
                    "破损退货",
                    "其中现场发现",
                    "冲销数量",
                    "净退货数量",
                ]
            )
            for row in summary[category]:
                sheet.append(
                    [
                        safe_cell(row["label"]),
                        safe_cell(row["base_unit"]),
                        row["confirmed_count"],
                        row["reversed_count"],
                        row["normal_qty"],
                        row["damaged_qty"],
                        row["new_damage_qty"],
                        row["reverse_qty"],
                        row["net_qty"],
                    ]
                )
        sheet = wb.create_sheet("退货明细")
        sheet.append(
            [
                "单号",
                "动作",
                "时间",
                "货主",
                "退货类型",
                "接收方",
                "操作员",
                "商品",
                "编码",
                "规格",
                "单位",
                "库位",
                "批次",
                "效期",
                "序列号",
                "正常",
                "已登记破损",
                "现场破损",
                "原因",
            ]
        )
        for event in events:
            line, order = event.line, event.line.order
            sheet.append(
                [
                    safe_cell(value)
                    for value in [
                        order.return_no,
                        event.get_kind_display(),
                        str(event.occurred_at),
                        order.owner.name,
                        order.get_recipient_type_display(),
                        order.recipient_name,
                        event.actor.username,
                        line.product_name,
                        line.product_code,
                        line.spec,
                        line.base_unit,
                        line.location_name,
                        line.dimensions["batch_no"],
                        line.dimensions["expiry_date"] or "",
                        line.dimensions["serial_no"],
                        line.available_qty,
                        line.damaged_qty,
                        line.new_damage_qty,
                        event.reason,
                    ]
                ]
            )
        output = BytesIO()
        wb.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="pda-returns.xlsx"'
        return response


def safe_cell(value):
    return (
        "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    )
