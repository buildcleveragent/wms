"""Real MySQL checks of physical returns, quality state, isolation and retry safety."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from allapp.accounts.models import UserRoleScope
from allapp.baseinfo.models import Owner, OwnerWarehouseBinding, Supplier
from allapp.inbound.models import Lot
from allapp.inventory.fifo import create_receipt_layer
from allapp.inventory.models import (
    InventoryDetail,
    InventoryLayerPosition,
    InventorySnapshotDaily,
    InventorySummary,
    InventoryTransaction,
)
from allapp.inventory.snapshot_services import generate_inventory_snapshot_for_date
from allapp.outbound.return_models import ReturnEvent, ReturnOrder
from allapp.products.identifier_services import add_product_barcode
from allapp.products.models import Product

BASE = "/api/pda/return-orders/"


@pytest.fixture
def env(db):
    from allapp.outbound.test_assisted_flow import AssistedOutboundFlowTests

    fixture = AssistedOutboundFlowTests()
    fixture.setUp()
    fixture.operator.user_permissions.add(
        Permission.objects.get(content_type__app_label="outbound", codename="process_pda_return")
    )
    fixture.client = APIClient()
    fixture.client.force_authenticate(fixture.operator)
    fixture.detail.onhand_qty = Decimal(20)
    fixture.detail.allocated_qty = Decimal(3)
    fixture.detail.locked_qty = Decimal(2)
    fixture.detail.damaged_qty = Decimal(5)
    fixture.detail.save()
    return fixture


def payload(env, **overrides):
    return {
        "request_id": str(uuid4()),
        "owner_id": env.owner.pk,
        "recipient_type": "OWNER",
        "reason": "破损退货",
        "lines": [
            {
                "inventory_id": env.detail.pk,
                "available_qty": "2",
                "damaged_qty": "3",
                "new_damage_qty": "1",
            }
        ],
        **overrides,
    }


def create(env, data=None):
    response = env.client.post(BASE, data or payload(env), format="json")
    assert response.status_code == 201, response.data
    return response.data


def action(env, row, name="confirm", **overrides):
    data = {"version": row["version"], "request_id": str(uuid4()), **overrides}
    if name == "reverse":
        data.setdefault("reason", "误退，货物尚未离库")
    return env.client.post(f"{BASE}{row['id']}/{name}/", data, format="json")


def test_mixed_return_and_reverse_preserves_new_damage(env):
    row = create(env)
    response = action(env, row)
    assert response.status_code == 200, response.data
    env.detail.refresh_from_db()
    assert (env.detail.onhand_qty, env.detail.available_qty, env.detail.damaged_qty) == (14, 7, 2)
    assert (env.detail.allocated_qty, env.detail.locked_qty) == (3, 2)
    assert InventoryTransaction.objects.get(src_model="ReturnOrder").qty_delta == -6
    response = action(env, response.data, "reverse")
    assert response.status_code == 200, response.data
    env.detail.refresh_from_db()
    assert (env.detail.onhand_qty, env.detail.available_qty, env.detail.damaged_qty) == (20, 9, 6)
    assert InventorySummary.objects.get(owner=env.owner, product=env.product).damaged_qty == 6
    assert ReturnEvent.objects.count() == 3


def test_idempotent_creation_confirm_reverse_and_deleted_request(env):
    data = payload(env)
    row = create(env, data)
    assert create(env, data)["id"] == row["id"]
    conflict = env.client.post(BASE, {**data, "reason": "不同内容"}, format="json")
    assert conflict.status_code == 409
    key = str(uuid4())
    done = action(env, row, request_id=key)
    assert done.status_code == 200
    assert action(env, row, request_id=key).status_code == 200
    assert action(env, row).status_code == 409
    reverse_key = str(uuid4())
    assert action(env, done.data, "reverse", request_id=reverse_key).status_code == 200
    assert action(env, done.data, "reverse", request_id=reverse_key).status_code == 200
    assert (
        action(env, done.data, "reverse", request_id=reverse_key, reason="改变原因").status_code
        == 409
    )
    assert InventoryTransaction.objects.filter(src_model="ReturnOrder").count() == 2
    draft_data = payload(env)
    draft = create(env, draft_data)
    assert (
        env.client.delete(f"{BASE}{draft['id']}/", {"version": 1}, format="json").status_code == 204
    )
    # Recovery after a lost response must resolve instead of leaving the PDA request pending.
    assert (
        env.client.delete(f"{BASE}{draft['id']}/", {"version": 1}, format="json").status_code == 204
    )
    assert (
        env.client.delete(f"{BASE}{draft['id']}/", {"version": 2}, format="json").status_code == 409
    )
    assert env.client.post(BASE, draft_data, format="json").status_code == 409
    assert ReturnOrder.all_objects.get(pk=draft["id"]).is_deleted


@pytest.mark.parametrize(
    "quantities",
    [
        ("11", "0", "0"),
        ("8", "0", "3"),
        ("0", "6", "0"),
        ("0", "0", "0"),
        ("-1", "0", "0"),
        ("0.00001", "0", "0"),
    ],
)
def test_quantity_rejections_never_touch_stock(env, quantities):
    a, d, n = quantities
    data = payload(
        env,
        lines=[
            {
                "inventory_id": env.detail.pk,
                "available_qty": a,
                "damaged_qty": d,
                "new_damage_qty": n,
            }
        ],
    )
    assert env.client.post(BASE, data, format="json").status_code == 400
    assert not ReturnOrder.objects.exists()
    assert not InventoryTransaction.objects.exists()


def test_stale_stock_and_post_failure_are_atomic(env):
    row = create(env)
    env.detail.locked_qty = 10
    env.detail.save()
    assert action(env, row).status_code == 409
    env.detail.locked_qty = 2
    env.detail.save()
    with patch(
        "allapp.outbound.return_services._refresh_summaries",
        side_effect=ValueError("summary failure"),
    ):
        assert action(env, row).status_code == 400
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == 20 and env.detail.damaged_qty == 5
    assert not InventoryTransaction.objects.exists() and not ReturnEvent.objects.exists()
    assert ReturnOrder.objects.get().status == "DRAFT"


def test_supplier_provenance_and_owner_fallback(env):
    supplier = Supplier.objects.create(owner=env.owner, name="原供应商")
    data = payload(env, recipient_type="SUPPLIER", supplier_id=supplier.pk)
    assert env.client.post(BASE, data, format="json").status_code == 400
    lot = Lot.objects.create(
        owner=env.owner,
        product=env.product,
        product_code=env.product.code,
        lot_no="RET-BATCH",
        supplier=supplier,
    )
    env.detail.lot = lot
    env.detail.batch_no = lot.lot_no
    env.detail.save()
    options = env.client.get(BASE + "stock-options/", {"product_id": env.product.pk})
    assert options.data["results"][0]["supplier_id"] == supplier.pk
    row = create(env, data)
    assert action(env, row).status_code == 200
    owner_draft = create(
        env, payload(env, lines=[{"inventory_id": env.detail.pk, "available_qty": "1"}])
    )
    assert owner_draft["recipient_type"] == "OWNER"


def test_shared_barcode_lists_owners_without_combining_stock(env):
    other = Owner.objects.create(code="RET-OTHER", name="另一货主")
    OwnerWarehouseBinding.objects.create(owner=other, warehouse=env.warehouse)
    product = Product.objects.create(
        owner=other, code="RETURN-SKU", name=env.product.name, base_uom=env.uom
    )
    detail = InventoryDetail.objects.create(
        owner=other, warehouse=env.warehouse, product=product, location=env.location, onhand_qty=7
    )
    for item in (env.product, product):
        add_product_barcode(product=item, barcode="RETURN-SHARED", barcode_type="UNIT")
    response = env.client.get(BASE + "products/", {"search": "RETURN-SHARED"})
    assert {row["owner_id"] for row in response.data["results"]} == {env.owner.pk, other.pk}
    bad = payload(env, lines=[{"inventory_id": detail.pk, "available_qty": "1"}])
    assert env.client.post(BASE, bad, format="json").status_code == 400
    data = payload(
        env, owner_id=other.pk, lines=[{"inventory_id": detail.pk, "available_qty": "1"}]
    )
    assert action(env, create(env, data)).status_code == 200
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == 20
    detail.refresh_from_db()
    assert detail.onhand_qty == 6
    OwnerWarehouseBinding.objects.filter(owner=other).update(is_active=False)
    assert env.client.get(BASE + "stock-options/", {"product_id": product.pk}).data["count"] == 0
    assert env.client.get(BASE).data["count"] == 0


def test_operator_cannot_edit_or_reverse_other_operators_order(env):
    row = create(env)
    done = action(env, row).data
    user = get_user_model().objects.create_user(username="return-other-op", warehouse=env.warehouse)
    UserRoleScope.objects.create(
        user=user, role=UserRoleScope.Role.WAREHOUSE_OPERATOR, warehouse=env.warehouse
    )
    user.user_permissions.add(Permission.objects.get(codename="process_pda_return"))
    env.client.force_authenticate(user)
    assert env.client.get(BASE).data["count"] == 1
    assert action(env, done, "reverse").status_code == 403
    env.client.force_authenticate(None)
    assert env.client.get(BASE).status_code in (401, 403)


def test_edit_versions_and_completed_readonly(env):
    row = create(env)
    data = payload(env, version=1, reason="修改原因")
    changed = env.client.patch(f"{BASE}{row['id']}/", data, format="json")
    assert changed.status_code == 200
    assert changed.data["version"] == 2
    assert action(env, row).status_code == 409
    done = action(env, changed.data)
    assert done.status_code == 200
    assert (
        env.client.patch(f"{BASE}{row['id']}/", {**data, "version": 3}, format="json").status_code
        == 409
    )


@override_settings(INVENTORY_FIFO_ENABLED=True)
def test_fifo_return_and_reverse_restore_original_positions(env):
    layer, _ = create_receipt_layer(
        owner_id=env.owner.pk,
        warehouse_id=env.warehouse.pk,
        product_id=env.product.pk,
        location_id=env.location.pk,
        quantity=20,
        source_type="TEST",
        source_id="return",
        unit_cost=Decimal(5),
        cost_currency="CNY",
    )
    done = action(env, create(env))
    assert done.status_code == 200, done.data
    assert InventoryLayerPosition.objects.get(layer=layer).remaining_qty == 14
    assert action(env, done.data, "reverse").status_code == 200
    assert InventoryLayerPosition.objects.get(layer=layer).remaining_qty == 20


@override_settings(INVENTORY_FIFO_ENABLED=True)
def test_missing_fifo_layer_rolls_back(env):
    row = create(env)
    assert action(env, row).status_code == 400
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == 20
    assert not InventoryTransaction.objects.exists()


def test_statistics_and_export_use_event_dates_and_units(env):
    done = action(env, create(env)).data
    yesterday = timezone.now() - timedelta(days=1)
    ReturnEvent.objects.filter(kind__in=["CONFIRM", "DAMAGE"]).update(occurred_at=yesterday)
    ReturnOrder.objects.filter(pk=done["id"]).update(confirmed_at=yesterday)
    assert action(env, done, "reverse").status_code == 200
    today_stats = env.client.get(BASE + "stats/")
    assert today_stats.status_code == 200
    assert today_stats.data["confirmed_count"] == 0 and today_stats.data["reversed_count"] == 1
    assert today_stats.data["units"][0]["net_qty"] == -6
    params = {
        "start_date": yesterday.date().isoformat(),
        "end_date": timezone.now().date().isoformat(),
    }
    stats = env.client.get(BASE + "stats/", params)
    assert stats.data["units"][0]["net_qty"] == 0
    assert stats.data["units"][0]["damaged_qty"] == 4
    response = env.client.get(BASE + "export/", params)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["退货明细"].max_row == 3
    assert env.client.get(BASE + "stats/", {"owner_id": "bad"}).status_code == 400


def test_quality_changes_roll_forward_in_daily_snapshot(env):
    now = timezone.now()
    yesterday = now.date() - timedelta(days=1)
    generate_inventory_snapshot_for_date(
        yesterday, owner_id=env.owner.pk, warehouse_id=env.warehouse.pk, bootstrap=True
    )
    confirmed = action(env, create(env))
    assert confirmed.status_code == 200
    generate_inventory_snapshot_for_date(
        now.date(), owner_id=env.owner.pk, warehouse_id=env.warehouse.pk
    )
    snapshot = InventorySnapshotDaily.objects.get(snapshot_date=now.date(), product=env.product)
    assert (snapshot.onhand_qty, snapshot.damaged_qty, snapshot.available_qty) == (14, 2, 7)
    assert action(env, confirmed.data, "reverse").status_code == 200
    generate_inventory_snapshot_for_date(
        now.date(), owner_id=env.owner.pk, warehouse_id=env.warehouse.pk
    )
    snapshot = InventorySnapshotDaily.objects.get(snapshot_date=now.date(), product=env.product)
    assert (snapshot.onhand_qty, snapshot.damaged_qty, snapshot.available_qty) == (20, 6, 9)


def test_multiline_return_is_atomic_and_counted_once(env):
    second = InventoryDetail.objects.create(
        owner=env.owner,
        warehouse=env.warehouse,
        product=env.product,
        location=env.location,
        batch_no="SECOND",
        onhand_qty=8,
    )
    row = create(
        env,
        payload(
            env,
            lines=[
                {"inventory_id": env.detail.pk, "available_qty": "2"},
                {"inventory_id": second.pk, "available_qty": "3"},
            ],
        ),
    )
    second.locked_qty = 6
    second.save()
    assert action(env, row).status_code == 409
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == 20
    assert not ReturnEvent.objects.exists() and not InventoryTransaction.objects.exists()
    second.locked_qty = 0
    second.save()
    assert action(env, row).status_code == 200
    stats = env.client.get(BASE + "stats/").data
    assert stats["confirmed_count"] == 1
    assert stats["units"][0]["normal_qty"] == 5


def test_revoked_permission_and_other_warehouse_cannot_read_or_write(env):
    from allapp.locations.models import Warehouse

    row = create(env)
    warehouse = Warehouse.objects.create(code="RET-DENY", name="Other warehouse")
    UserRoleScope.objects.filter(user=env.operator).update(warehouse=warehouse)
    assert env.client.get(f"{BASE}{row['id']}/").status_code == 404
    assert env.client.get(BASE + "products/").data["count"] == 0
    assert action(env, row).status_code == 404
    UserRoleScope.objects.filter(user=env.operator).update(warehouse=env.warehouse)
    env.operator.user_permissions.clear()
    env.client.force_authenticate(get_user_model().objects.get(pk=env.operator.pk))
    assert env.client.get(BASE).status_code == 403
    assert not InventoryTransaction.objects.exists()


def test_soft_deleted_products_excluded_and_source_change_rejected(env):
    row = create(env)
    Product.objects.filter(pk=env.product.pk).update(is_deleted=True)
    assert env.client.get(BASE + "products/").data["count"] == 0
    assert action(env, row).status_code == 409
    Product.all_objects.filter(pk=env.product.pk).update(is_deleted=False)
    env.detail.batch_no = "CHANGED"
    env.detail.save()
    assert action(env, row).status_code == 409


@pytest.mark.django_db(transaction=True)
@pytest.mark.parametrize("operation", ["competing", "same_confirm", "same_reverse"])
def test_concurrent_return_writes_are_exactly_once(env, operation):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    from django.db import close_old_connections

    data = payload(env, lines=[{"inventory_id": env.detail.pk, "available_qty": "8"}])
    first = create(env, data)
    second = (
        create(env, {**data, "request_id": str(uuid4())}) if operation == "competing" else first
    )
    if operation == "same_reverse":
        first = second = action(env, first).data
    barrier = Barrier(2)
    key = str(uuid4())
    user_id = env.operator.pk

    def submit(row):
        close_old_connections()
        try:
            client = APIClient()
            client.force_authenticate(get_user_model().objects.get(pk=user_id))
            barrier.wait(timeout=10)
            name = "reverse" if operation == "same_reverse" else "confirm"
            return client.post(
                f"{BASE}{row['id']}/{name}/",
                {"version": row["version"], "request_id": key, "reason": "并发测试"},
                format="json",
            ).status_code
        finally:
            close_old_connections()

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [executor.submit(submit, row) for row in (first, second)]
        codes = sorted(future.result(timeout=30) for future in futures)
    assert codes == ([200, 409] if operation == "competing" else [200, 200])
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == (20 if operation == "same_reverse" else 12)
    assert InventoryTransaction.objects.filter(src_model="ReturnOrder").count() == (
        2 if operation == "same_reverse" else 1
    )


def test_serial_return_and_fifo_configuration_change_fail_closed(env):
    env.product.serial_control = True
    env.product.save()
    env.detail.serial_no = "RET-SN1"
    env.detail.onhand_qty = 1
    env.detail.allocated_qty = env.detail.locked_qty = 0
    env.detail.damaged_qty = 1
    env.detail.save()
    data = payload(env, lines=[{"inventory_id": env.detail.pk, "damaged_qty": "1"}])
    done = action(env, create(env, data))
    assert done.status_code == 200
    with override_settings(INVENTORY_FIFO_ENABLED=True):
        assert action(env, done.data, "reverse").status_code == 400
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == 0
    assert action(env, done.data, "reverse").status_code == 200
    env.detail.refresh_from_db()
    assert env.detail.onhand_qty == env.detail.damaged_qty == 1


def test_deactivated_origin_blocks_reversal_without_partial_restore(env):
    done = action(env, create(env)).data
    InventoryDetail.objects.filter(pk=env.detail.pk).update(is_active=False)
    assert action(env, done, "reverse").status_code == 409
    assert ReturnOrder.objects.get(pk=done["id"]).status == "COMPLETED"
    assert InventoryTransaction.objects.count() == 1


def test_wire_quantities_are_decimal_strings_and_export_escapes_formulas(env):
    response = env.client.get(BASE + "products/")
    assert isinstance(response.json()["results"][0]["available_qty"], str)
    from allapp.outbound.return_api import safe_cell

    assert safe_cell('=HYPERLINK("example")').startswith("'=")
