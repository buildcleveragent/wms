"""Real MySQL export contracts: exact rows, private scope and no business writes."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import resolve
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from allapp.accounts.models import AuditEvent, UserRoleScope
from allapp.baseinfo.models import Customer
from allapp.core import owner_exports
from allapp.outbound.models import OutboundOrder
from allapp.outbound.views import OutboundOrderViewSet
from allapp.products.models import ProductPackage

from . import test_owner_detail as fixtures
from .models import InventoryDetail, InventorySummary
from .views import OwnerInventorySummaryViewSet


class OwnerExportTests(TestCase):
    setUpTestData = classmethod(fixtures.OwnerInventoryDetailTests.setUpTestData.__func__)
    stock = fixtures.OwnerInventoryDetailTests.stock

    def setUp(self):
        self.order = OutboundOrder.objects.create(
            order_no="0000123",
            owner=self.owner,
            warehouse=self.warehouse,
            customer=self.customer,
            created_by=self.user,
            biz_date=date(2026, 10, 19),
        )
        self.line = self.order.lines.create(
            product=self.product,
            base_qty=Decimal("2.125"),
            base_price=Decimal("3.1234"),
            note="=1+1",
        )

    def call(self, action="owner_export", params=None, user=None, pk=None):
        request = APIRequestFactory().get("/", params or {})
        force_authenticate(request, user=user or self.user)
        cls = (
            OutboundOrderViewSet
            if action in {"owner_export", "list"}
            else OwnerInventorySummaryViewSet
        )
        return cls.as_view({"get": action})(request, **({"pk": pk} if pk else {}))

    def workbook(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, "data", ""))
        self.assertIn("private", response["Cache-Control"])
        try:
            return load_workbook(BytesIO(b"".join(response.streaming_content)))
        finally:
            response.file_to_stream.close()

    def test_orders_exact_lines_amounts_text_and_pagination_ignored(self):
        package = ProductPackage.objects.create(product=self.product, uom=self.unit, qty_in_base=1)
        self.line.aux_uom = package
        self.line.aux_qty = Decimal("2.125")
        self.line.save()
        self.order.lines.create(product=self.product, base_qty=1, base_price=3)
        wb = self.workbook(self.call(params={"search": "0000123", "page": 999, "page_size": 1}))
        sheet = wb["订单商品明细"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["A2"].value, "0000123")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["Q2"].value, 6.64)
        self.assertEqual(sheet["R2"].value, "瓶")
        self.assertEqual(sheet["T2"].value, "=1+1")
        self.assertEqual(sheet["T2"].data_type, "s")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:U3")
        self.assertTrue(AuditEvent.objects.filter(action="EXPORT", module="owner.orders").exists())

    def test_deleted_lines_leave_placeholder_deleted_orders_excluded(self):
        self.order.lines.update(is_deleted=True)
        sheet = self.workbook(self.call())["订单商品明细"]
        self.assertEqual(sheet["L2"].value, "无有效商品行")
        self.assertIsNone(sheet["O2"].value)
        self.assertIsNone(sheet["Q2"].value)
        OutboundOrder.objects.filter(pk=self.order.pk).update(is_deleted=True)
        self.assertEqual(self.call().status_code, 400)

    @override_settings(OUTBOUND_LEGACY_AUTHZ_MODE="shadow")
    def test_salesperson_list_and_export_strict_even_in_legacy_mode(self):
        other = get_user_model().objects.create_user(username="export-other")
        OutboundOrder.objects.create(
            owner=self.owner, warehouse=self.warehouse, customer=self.customer, created_by=other
        )
        OutboundOrder.objects.create(
            owner=self.other,
            warehouse=self.warehouse,
            customer=Customer.objects.create(
                owner=self.other, salesperson=self.user, code="EX-OTHER", name="其他客户"
            ),
            created_by=self.user,
        )
        response = self.call("list", {"owner_portal": "1", "search": "000"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([r["id"] for r in response.data["results"]], [self.order.pk])
        self.assertEqual(self.workbook(self.call())["订单商品明细"].max_row, 2)

    def test_manager_sees_same_owner_orders_only(self):
        manager = get_user_model().objects.create_user(username="export-manager")
        UserRoleScope.objects.create(user=manager, role="owner_manager", owner=self.owner)
        manager.user_permissions.add(Permission.objects.get(codename="view_outboundorder"))
        OutboundOrder.objects.create(
            owner=self.owner, warehouse=self.warehouse, customer=self.customer, created_by=manager
        )
        OutboundOrder.objects.create(
            owner=self.other,
            warehouse=self.warehouse,
            customer=Customer.objects.create(
                owner=self.other, salesperson=self.user, code="EX-OTHER", name="其他客户"
            ),
            created_by=manager,
        )
        self.assertEqual(self.workbook(self.call(user=manager))["订单商品明细"].max_row, 3)

    def test_permission_revocation_and_warehouse_role(self):
        self.user.user_permissions.clear()
        fresh = get_user_model().objects.get(pk=self.user.pk)
        self.assertEqual(self.call(user=fresh).status_code, 403)
        wh = get_user_model().objects.create_user(username="export-warehouse")
        UserRoleScope.objects.create(user=wh, role="warehouse_manager", warehouse=self.warehouse)
        for action in ("owner_export", "export_summary", "export_detail"):
            self.assertEqual(
                self.call(
                    action, user=wh, pk=self.summary.pk if action == "export_detail" else None
                ).status_code,
                403,
            )

    def test_empty_invalid_and_cross_owner_inputs(self):
        for params in (
            {"search": "不存在"},
            {"warehouse_id": "bad"},
            {"date_from": "bad"},
            {"unexpected": "1"},
        ):
            self.assertEqual(self.call(params=params).status_code, 400, params)
        self.assertEqual(self.call(params={"owner_id": self.other.pk}).status_code, 403)
        foreign = InventorySummary.objects.create(owner=self.other, product=self.product)
        self.assertEqual(self.call("export_detail", pk=foreign.pk).status_code, 404)

    def test_summary_raw_quantities_and_no_rebuild(self):
        InventorySummary.objects.filter(pk=self.summary.pk).update(
            onhand_qty="12.2500",
            allocated_qty=1,
            locked_qty=2,
            damaged_qty=3,
            available_qty="6.2500",
        )
        sheet = self.workbook(self.call("export_summary", {"search": "库存", "page": 999}))[
            "商品库存汇总"
        ]
        self.assertEqual([sheet.cell(2, i).value for i in range(6, 11)], [12.25, 6.25, 1, 2, 3])
        self.assertEqual(InventoryDetail.objects.count(), 0)

    def test_batches_merge_locations_and_keep_units_separate(self):
        self.stock("2.125", batch_no="00001")
        self.stock("3.25", location=self.location2, batch_no="00001")
        other_unit = self.stock("1", location=self.remote, batch_no="00001")
        InventoryDetail.objects.filter(pk=other_unit.pk).update(base_unit="历史单位")
        self.stock("0", batch_no="ZERO")
        sheet = self.workbook(self.call("export_detail", pk=self.summary.pk))["商品批次库存明细"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["F2"].value, "00001")
        self.assertEqual(sheet["J2"].value, 5.375)
        self.assertEqual(sheet["I3"].value, "历史单位")
        self.assertEqual(sheet["G2"].value, "未记录")

    def test_expiry_and_warehouse_filters_match_detail(self):
        self.stock("2", batch_no="TODAY", expiry_date=date(2026, 10, 19))
        self.stock("3", batch_no="OLD", expiry_date=date(2026, 10, 18))
        self.stock("4", location=self.remote, batch_no="UNKNOWN")
        with patch(
            "allapp.inventory.owner_detail.timezone.now",
            return_value=__import__("datetime").datetime(2026, 10, 19, 12),
        ):
            for expiry, expected in [("valid", 2), ("expired", 3), ("unknown", 4)]:
                response = self.call("export_detail", {"expiry": expiry}, pk=self.summary.pk)
                sheet = self.workbook(response)["商品批次库存明细"]
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet["J2"].value, expected)
        self.assertEqual(
            self.call(
                "export_detail",
                {"warehouse_id": self.warehouse2.pk, "batch": "TODAY"},
                pk=self.summary.pk,
            ).status_code,
            400,
        )

    def test_small_limit_rejects_whole_export_and_read_has_no_business_writes(self):
        self.order.lines.create(product=self.product, base_qty=1, base_price=1)
        with patch.object(owner_exports, "MAX_EXPORT_ROWS", 1):
            response = self.call()
        self.assertEqual(response.status_code, 413)
        self.assertNotIn("Content-Disposition", response)
        with CaptureQueriesContext(connection) as captured:
            self.workbook(self.call())
        writes = [
            q["sql"]
            for q in captured
            if q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
        ]
        self.assertTrue(writes)
        self.assertTrue(all("accounts_auditevent" in sql for sql in writes), writes)

    def test_query_count_does_not_grow_per_line(self):
        with CaptureQueriesContext(connection) as first:
            self.workbook(self.call())
        for _ in range(12):
            self.order.lines.create(product=self.product, base_qty=1, base_price=1)
        with CaptureQueriesContext(connection) as second:
            self.workbook(self.call())
        self.assertLessEqual(len(second), len(first) + 1)

    def test_decimal_precision_and_exact_50000_boundary(self):
        request = APIRequestFactory().get("/")
        request.user = self.user
        kwargs = dict(
            request=request,
            scope=owner_exports.owner_scope(self.user),
            kind="inventory",
            title="边界测试",
            headers=["数量"],
            filters="",
            basis="测试",
        )
        rows = [[Decimal("123456789012345.1234")]] * 50000
        wb = self.workbook(owner_exports.xlsx_download(rows=rows, **kwargs))
        self.assertEqual(wb["边界测试"].max_row, 50001)
        self.assertEqual(wb["边界测试"]["A2"].value, "123456789012345.1234")
        with self.assertRaises(owner_exports.ExportTooLarge):
            owner_exports.xlsx_download(rows=rows + rows[:1], **kwargs)

    def test_audit_failure_does_not_return_file(self):
        with patch(
            "allapp.core.owner_exports.record_audit_event",
            side_effect=RuntimeError("audit unavailable"),
        ):
            with self.assertRaises(RuntimeError):
                self.call()

    def test_routes(self):
        for path in [
            "/api/outbound/orders/owner-export/",
            "/api/inventory/summary/export/",
            f"/api/inventory/summary/{self.summary.pk}/detail/export/",
        ]:
            self.assertIn("get", resolve(path).func.actions)

    def test_export_matches_all_list_pages_and_preserves_closed_orders(self):
        for index in range(2):
            OutboundOrder.objects.create(
                order_no=f"EX-PAGE-{index}",
                owner=self.owner,
                warehouse=self.warehouse,
                customer=self.customer,
                created_by=self.user,
                is_closed=bool(index),
                close_reason="测试关闭" if index else "",
                biz_date=date(2026, 10, 18 - index),
            )
        numbers = []
        for page in range(1, 4):
            result = self.call(
                "list", {"owner_portal": "1", "search": "保密客户", "page": page, "page_size": 1}
            )
            self.assertEqual(result.status_code, 200)
            numbers.extend(row["order_no"] for row in result.data["results"])
        sheet = self.workbook(self.call(params={"search": "保密客户"}))["订单商品明细"]
        self.assertEqual([row[0] for row in list(sheet.values)[1:]], numbers)
        self.assertEqual(sheet["H4"].value, "已关闭")

    def test_invalid_dates_and_scope_revocation(self):
        for params in (
            {"date_from": "2026-02-30"},
            {"date_from": "2026-10-20", "date_to": "2026-10-19"},
        ):
            self.assertEqual(self.call(params=params).status_code, 400)
        UserRoleScope.objects.filter(user=self.user).update(is_active=False)
        for action in ("owner_export", "export_summary"):
            self.assertEqual(self.call(action).status_code, 403)

    def test_actual_read_limit_rechecks_growth_after_count(self):
        qs = MagicMock()
        qs.count.return_value = 1
        qs.__getitem__.return_value = [{"id": 1}, {"id": 2}]
        with patch.object(owner_exports, "MAX_EXPORT_ROWS", 1):
            with self.assertRaises(owner_exports.ExportTooLarge):
                owner_exports.bounded_rows(qs)
