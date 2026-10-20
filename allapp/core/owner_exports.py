"""Bounded, private Excel downloads for an explicitly scoped owner account."""

from datetime import date, datetime
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from zoneinfo import ZoneInfo

from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import APIException, PermissionDenied, ValidationError

from allapp.accounts.access import AccessScope
from allapp.accounts.audit import record_audit_event
from allapp.baseinfo.models import Owner

MAX_EXPORT_ROWS = 50000
BUSINESS_TZ = ZoneInfo("Asia/Shanghai")


class ExportTooLarge(APIException):
    status_code = 413
    default_detail = "导出最多支持 50,000 行，请缩小筛选范围后重试。"


def owner_scope(user):
    scope = AccessScope.for_user(user)
    if (
        not scope.is_valid
        or scope.is_global
        or len(scope.owner_ids) != 1
        or not set(scope.roles).intersection({"owner_manager", "owner_salesperson"})
        or scope.warehouse_ids
    ):
        raise PermissionDenied("需要有效的单一货主角色范围。")
    return scope


def export_params(params, allowed):
    """Pagination never changes export scope; reject unsupported business filters."""
    ignored = {"page", "page_size", "cursor", "owner_portal", "format"}
    invalid = set(params) - set(allowed) - ignored
    if invalid:
        raise ValidationError("包含不支持的筛选条件。")
    return {k: params[k] for k in allowed if params.get(k) not in (None, "")}


def bounded_rows(qs):
    if qs.count() > MAX_EXPORT_ROWS:
        raise ExportTooLarge()
    rows = list(qs[: MAX_EXPORT_ROWS + 1])
    if len(rows) > MAX_EXPORT_ROWS:
        raise ExportTooLarge()
    if not rows:
        raise ValidationError("当前筛选下没有可导出的数据。")
    return rows


def _cell(sheet, value):
    cell = WriteOnlyCell(sheet)
    if isinstance(value, Decimal):
        # Excel only preserves 15 significant decimal digits.
        normalized = value.normalize()
        if len(normalized.as_tuple().digits) > 15 or normalized.adjusted() >= 15:
            value = format(value, "f")
        else:
            cell.value = value
            places = max(0, -value.as_tuple().exponent)
            cell.number_format = "0" + ("." + "0" * places if places else "")
            return cell
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"
        cell.number_format = "@"
    else:
        cell.value = value
    return cell


def xlsx_download(request, *, scope, kind, title, headers, rows, filters, basis):
    """Finish a complete workbook and audit before exposing any bytes to the client."""
    if len(rows) > MAX_EXPORT_ROWS:
        raise ExportTooLarge()
    if not rows:
        raise ValidationError("当前筛选下没有可导出的数据。")
    now = timezone.now()
    now = (now.replace(tzinfo=BUSINESS_TZ) if timezone.is_naive(now) else now).astimezone(
        BUSINESS_TZ
    )
    owner = Owner.objects.get(pk=scope.single_owner_id)
    workbook = Workbook(write_only=True)
    info = workbook.create_sheet("导出说明")
    scope_label = (
        "本人创建的订单"
        if kind == "orders" and "owner_salesperson" in scope.roles
        else "本货主共享库存" if kind != "orders" else "本货主订单"
    )
    for row in [
        ("项目", "内容"),
        ("报告", title),
        ("货主", owner.name),
        ("授权范围", scope_label),
        ("已应用筛选", filters or "未设置筛选"),
        ("导出时间（Asia/Shanghai）", now.isoformat()),
        ("数据行数", len(rows)),
        ("口径", basis),
        ("精度说明", "超过 Excel 安全数值精度的数量及金额以文本保留；编号、SKU、批次为文本。"),
    ]:
        info.append([_cell(info, v) for v in row])
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    sheet.append([_cell(sheet, v) for v in headers])
    for row in rows:
        sheet.append([_cell(sheet, v) for v in row])
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, suffix=".xlsx")
    try:
        workbook.save(output)
        record_audit_event(
            action="EXPORT",
            module=f"owner.{kind}",
            request=request,
            owner_id=scope.single_owner_id,
            metadata={"rows": len(rows), "filters": filters, "scope": scope_label},
        )
        output.seek(0)
        response = FileResponse(
            output,
            as_attachment=True,
            filename=f"owner-{kind}-{now:%Y%m%d-%H%M%S}.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Cache-Control"] = "private, no-store"
        response["X-Content-Type-Options"] = "nosniff"
        return response
    except BaseException:
        output.close()
        raise
