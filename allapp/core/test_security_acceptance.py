import io
import zipfile
from unittest.mock import patch

from django.contrib import admin
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.test import Client, RequestFactory, SimpleTestCase, TestCase
from openpyxl import load_workbook
from rest_framework_simplejwt.tokens import AccessToken

from allapp.accounts.models import UserRoleScope
from allapp.baseinfo.models import Customer
from allapp.inbound import tests as inbound_fixtures
from allapp.inbound.excel_import import build_no_order_receive_template
from allapp.outbound.models import OutboundOrder, OutboundOrderLine
from allapp.pos import tests as pos_fixtures
from allapp.products.identifier_excel import build_identifier_export
from allapp.products.identifier_services import add_external_identifier
from allapp.tasking.models import (
    TaskAssignment,
    TaskScanLog,
    TaskStatusLog,
    WmsTask,
    WmsTaskLine,
)

from .print_tickets import issue_print_ticket
from .xlsx_safety import validate_xlsx_archive


class XlsxBudgetTests(SimpleTestCase):
    def test_rejects_excessive_expansion_before_decompressing(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("sheet.xml", "a" * 100)
        with patch("allapp.core.xlsx_safety.MAX_EXPANDED_BYTES", 99):
            with patch.object(zipfile.ZipFile, "testzip") as decompress:
                with self.assertRaises(ValueError):
                    validate_xlsx_archive(stream.getvalue())
                decompress.assert_not_called()

    def test_rejects_too_many_entries_and_corrupt_archives(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("one", "ok")
            archive.writestr("two", "ok")
        with patch("allapp.core.xlsx_safety.MAX_ENTRIES", 1):
            with self.assertRaises(ValueError):
                validate_xlsx_archive(stream.getvalue())
        with self.assertRaises(ValueError):
            validate_xlsx_archive(b"not an xlsx")


class PosCatalogPermissionTests(TestCase):
    setUp = pos_fixtures.PosApiTests.setUp
    open_shift = pos_fixtures.PosApiTests.open_shift

    def test_authenticated_buyer_without_pos_permission_cannot_read_catalog(self):
        buyer = get_user_model().objects.create_user(
            username="scope-buyer", warehouse=self.warehouse, owner=self.owner
        )
        self.client.force_authenticate(buyer)
        for params in ({}, {"search": self.product.code}, {"barcode": "POS-CTN-BAR"}):
            response = self.client.get("/api/pos/products/", params)
            self.assertEqual(response.status_code, 403)

    def test_authorized_cashier_keeps_cross_owner_stock_visibility(self):
        response = self.client.get("/api/pos/products/", {"barcode": "POS-OTH-UNIT-BAR"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["results"][0]["id"], self.other_product.id)

    def test_explicit_warehouse_scope_revokes_stale_pos_binding(self):
        from allapp.pos.services import create_pos_sale

        other_warehouse = type(self.warehouse).objects.create(code="POSOTH", name="Other")

        UserRoleScope.objects.create(
            user=self.user,
            role=UserRoleScope.Role.WAREHOUSE_OPERATOR,
            warehouse=other_warehouse,
        )
        self.assertEqual(self.client.get("/api/pos/products/").status_code, 403)
        with self.assertRaises(ValidationError):
            create_pos_sale(user=self.user, items=[])


class SuperuserPasswordBoundaryTests(TestCase):
    def setUp(self):
        self.editor = get_user_model().objects.create_user(
            username="password-editor", is_staff=True
        )
        self.editor.user_permissions.add(
            Permission.objects.get(content_type__app_label="accounts", codename="change_user")
        )
        self.target = get_user_model().objects.create_superuser(
            username="password-target",
            email="target@example.test",
            password="test-original",
        )
        self.client.force_login(self.editor)

    def test_delegated_editor_cannot_open_or_submit_superuser_password_form(self):
        original = self.target.password
        url = f"/admin/accounts/user/{self.target.pk}/password/"
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(
            self.client.post(
                url,
                {
                    "password1": "Changed-test-password-345!",
                    "password2": "Changed-test-password-345!",
                },
            ).status_code,
            403,
        )
        self.target.refresh_from_db()
        self.assertEqual(self.target.password, original)

    def test_delegated_editor_can_still_reset_an_ordinary_user(self):
        ordinary = get_user_model().objects.create_user(username="ordinary-target")
        response = self.client.post(
            f"/admin/accounts/user/{ordinary.pk}/password/",
            {
                "password1": "Changed-test-password-345!",
                "password2": "Changed-test-password-345!",
            },
        )
        self.assertEqual(response.status_code, 302)
        ordinary.refresh_from_db()
        self.assertTrue(ordinary.check_password("Changed-test-password-345!"))


class TaskAndPrintBoundaryTests(TestCase):
    setUp = inbound_fixtures.InboundAuthorizationAndWorkflowTests.setUp

    def _manager(self):
        user = get_user_model().objects.create_user(username="scoped-manager", is_staff=True)
        UserRoleScope.objects.create(
            user=user, role=UserRoleScope.Role.WAREHOUSE_MANAGER, warehouse=self.warehouse
        )
        user.user_permissions.add(
            *Permission.objects.filter(
                content_type__app_label="tasking",
                codename__in=[
                    "view_wmstask",
                    "view_wmstaskline",
                    "change_wmstaskline",
                    "view_taskassignment",
                    "change_taskassignment",
                    "taskconfirm_as_wh_manager",
                ],
            )
        )
        return user

    def _task(self, code, other=False, kind="RECEIVE"):
        return WmsTask.objects.create(
            task_no=code,
            task_type=kind,
            owner=self.other_owner if other else self.owner,
            warehouse=self.other_warehouse if other else self.warehouse,
        )

    def test_task_log_admin_querysets_are_scoped(self):
        user = self._manager()
        own, other = self._task("LOG-OWN"), self._task("LOG-OTHER", other=True)
        request = RequestFactory().get("/admin/")
        request.user = user
        for model in (TaskStatusLog, TaskScanLog):
            rows = []
            for task in (own, other):
                values = {"task": task}
                if model is TaskStatusLog:
                    values.update(old_status="RELEASED", new_status="IN_PROGRESS")
                else:
                    values.update(
                        owner=task.owner,
                        warehouse=task.warehouse,
                        fp=f"log-{task.pk}",
                        scan_snapshot_rev=1,
                        qty_base_delta=1,
                    )
                rows.append(model.objects.create(**values))
            model_admin = admin.site._registry[model]
            self.assertEqual(list(model_admin.get_queryset(request)), rows[:1])
            self.assertFalse(model_admin.has_change_permission(request, rows[1]))

    def test_inbound_template_treats_product_names_as_text(self):
        self.product.name = "=1+1"
        self.product.save(update_fields=["name"])
        workbook = load_workbook(io.BytesIO(build_no_order_receive_template(self.owner)))
        matching = [
            cell for sheet in workbook for row in sheet for cell in row if cell.value == "=1+1"
        ]
        self.assertTrue(matching)
        self.assertTrue(all(cell.data_type == "s" for cell in matching))

    def test_task_children_and_foreign_keys_respect_manager_scope(self):
        user = self._manager()
        own, other = self._task("OWN"), self._task("OTHER", other=True)
        own_line = WmsTaskLine.objects.create(task=own, product=self.product, qty_plan=1)
        other_product = type(self.product).objects.create(
            owner=self.other_owner,
            code="OTHER-ADMIN",
            sku="OTHER-ADMIN",
            name="Other product",
            base_uom=self.base_uom,
        )
        other_line = WmsTaskLine.objects.create(task=other, product=other_product, qty_plan=1)
        own_assignment = TaskAssignment.objects.create(task=own, assignee=user)
        TaskAssignment.objects.create(task=other, assignee=user)
        request = RequestFactory().get("/admin/")
        request.user = user
        for model, expected in ((WmsTaskLine, own_line), (TaskAssignment, own_assignment)):
            model_admin = admin.site._registry[model]
            self.assertEqual(list(model_admin.get_queryset(request)), [expected])
            field = model_admin.formfield_for_foreignkey(model._meta.get_field("task"), request)
            self.assertEqual(list(field.queryset), [own])
            with self.assertRaises(ValidationError):
                model_admin._validate_values(request, {"task": other})
        self.assertFalse(
            admin.site._registry[WmsTaskLine].has_change_permission(request, other_line)
        )
        with self.assertRaises(ValidationError):
            admin.site._registry[TaskAssignment]._validate_values(
                request, {"task": own, "line": other_line}
            )

    def test_print_tokens_are_revoked_by_password_change(self):
        user = self._manager()
        token = str(AccessToken.for_user(user))
        tasks = [self._task("PRINT-RECEIVE"), self._task("PRINT-PICK", kind="PICK")]
        order = OutboundOrder.objects.create(
            owner=self.owner,
            warehouse=self.warehouse,
            order_no="PRINT-ORDER",
            customer=Customer.objects.create(
                owner=self.owner, code="PRINT", name="Print", salesperson=user
            ),
            outbound_type="SALES",
        )
        line = OutboundOrderLine.objects.create(
            order=order,
            product=self.product,
            base_qty=1,
            base_price=0,
            base_uom=self.base_uom,
            line_no=1,
        )
        tasks[1].source_model = "OutboundOrder"
        tasks[1].source_pk = str(order.pk)
        tasks[1].save(update_fields=["source_model", "source_pk"])
        WmsTaskLine.objects.create(
            task=tasks[1],
            product=self.product,
            from_location=self.location,
            qty_plan=1,
            src_model="OutboundOrderLine",
            src_id=line.pk,
        )
        urls = [
            f"/api/inbound/receive_task/{tasks[0].pk}/print/",
            f"/api/pda/pick-tasks/{tasks[1].pk}/print/",
        ]
        client = Client()
        tickets = {url: issue_print_ticket(user, url) for url in urls}
        for url in urls:
            self.assertEqual(
                client.get(url, {"ticket": tickets[url], "kind": "pick"}).status_code, 200
            )
            self.assertEqual(client.get(url, {"token": token}).status_code, 401)
        user.set_password("Changed-print-test-password!")
        user.save(update_fields=["password"])
        for url in urls:
            self.assertEqual(
                client.get(url, {"ticket": tickets[url], "kind": "pick"}).status_code, 401
            )
            self.assertEqual(client.get(url, {"token": token}).status_code, 401)
            self.assertEqual(client.get(url, {"access": token}).status_code, 401)

    def test_print_ticket_is_bound_to_document_and_cannot_authenticate_api(self):
        user = self._manager()
        own = self._task("TICKET-OWN")
        other = self._task("TICKET-OTHER", other=True)
        path = f"/api/inbound/receive_task/{own.pk}/print/"
        ticket = issue_print_ticket(user, path)
        client = Client()
        self.assertEqual(
            client.get(
                f"/api/inbound/receive_task/{other.pk}/print/", {"ticket": ticket}
            ).status_code,
            401,
        )
        self.assertEqual(client.get(path, {"ticket": ticket + "x"}).status_code, 401)
        with patch("django.core.signing.time.time", return_value=__import__("time").time() + 121):
            self.assertEqual(client.get(path, {"ticket": ticket}).status_code, 401)
        self.assertEqual(
            client.get("/api/auth/profile/", HTTP_AUTHORIZATION=f"Bearer {ticket}").status_code,
            401,
        )
        user.is_active = False
        user.save(update_fields=["is_active"])
        self.assertEqual(client.get(path, {"ticket": ticket}).status_code, 401)

    def test_external_identifiers_export_as_literal_text(self):
        value = '=HYPERLINK("https://example.invalid","identifier")'
        add_external_identifier(product=self.product, source_system="TEST", external_code=value)
        workbook = load_workbook(io.BytesIO(build_identifier_export(self.owner)), data_only=False)
        cell = workbook["外部标识维护"]["E2"]
        self.assertEqual(cell.value, value)
        self.assertEqual(cell.data_type, "s")
