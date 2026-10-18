"""Boss dashboard input and export boundaries against the real API."""

from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from allapp.accounts.models import UserRoleScope
from allapp.accounts.roles import ROLE_GROUP_TEMPLATES
from allapp.locations.models import Warehouse
from allapp.reports.models import BusinessReviewSnapshot


class BossInputAndExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("sync_wms_role_groups", stdout=StringIO())
        cls.warehouse = Warehouse.objects.create(code="BOSS-REG-A", name="Boss A")
        cls.other_warehouse = Warehouse.objects.create(code="BOSS-REG-B", name="Boss B")
        cls.user = get_user_model().objects.create_user(username="boss-regression")
        cls.group = Group.objects.get(
            name=ROLE_GROUP_TEMPLATES[UserRoleScope.Role.WAREHOUSE_BOSS].group_name
        )
        UserRoleScope.objects.create(
            user=cls.user, role=UserRoleScope.Role.WAREHOUSE_BOSS, warehouse=cls.warehouse
        )
        cls.user.groups.add(cls.group)
        now = timezone.now()
        today = timezone.localdate(now) if timezone.is_aware(now) else now.date()
        cls.snapshot = BusinessReviewSnapshot.objects.create(
            share_code="boss-regression-snapshot",
            created_by=cls.user,
            warehouse_ids=[cls.warehouse.pk],
            date_from=today,
            date_to=today,
            data_as_of=now,
            payload={},
            checksum="0" * 64,
        )
        cls.other_snapshot = BusinessReviewSnapshot.objects.create(
            share_code="boss-regression-other",
            created_by=cls.user,
            warehouse_ids=[cls.other_warehouse.pk],
            date_from=today,
            date_to=today,
            data_as_of=now,
            payload={},
            checksum="0" * 64,
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def revoke(self, codename):
        permission = Permission.objects.get(content_type__app_label="reports", codename=codename)
        self.group.permissions.remove(permission)
        self.user.user_permissions.remove(permission)
        for key in ("_perm_cache", "_user_perm_cache", "_group_perm_cache"):
            self.user.__dict__.pop(key, None)

    def test_invalid_end_date_is_400_on_every_dated_boss_endpoint(self):
        endpoints = (
            "home",
            "inventory",
            "inventory/details",
            "alerts",
            "revenue-assurance",
            "receivables",
            "receivables/bills",
            "operations",
            "operations/details",
            "resource-yield",
            "performance",
            "inventory-risk",
            "inventory-risk/details",
            "alert-cases",
        )
        self.client.raise_request_exception = False
        for endpoint in endpoints:
            with self.subTest(endpoint=endpoint):
                response = self.client.get(f"/api/reports/boss/{endpoint}/", {"date_to": "bad"})
                self.assertEqual(response.status_code, 400)
                self.assertIn("detail", response.data)

    def test_invalid_post_dates_reject_before_export_or_snapshot_creation(self):
        self.client.raise_request_exception = False
        for endpoint in ("exports", "review-snapshots"):
            for value in ("bad", "2026-02-30", [], {"date": "2026-01-01"}):
                with self.subTest(endpoint=endpoint, value=value):
                    response = self.client.post(
                        f"/api/reports/boss/{endpoint}/",
                        {"date_to": value, "report_type": "receivables"},
                        format="json",
                    )
                    self.assertEqual(response.status_code, 400)

    def test_export_revocation_blocks_both_downloads_before_workbook_generation(self):
        self.revoke("export_operations")
        # The legacy permission cannot substitute for the canonical report permission.
        self.user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="accounts", codename="export_operational_reports"
            )
        )
        with patch("allapp.reports.views_boss_p1._cockpit_workbook") as workbook:
            for report in (
                "receivables",
                "revenue_assurance",
                "performance",
                "resource_yield",
                "inventory_risk",
            ):
                with self.subTest(report=report):
                    response = self.client.post(
                        "/api/reports/boss/exports/", {"report_type": report}, format="json"
                    )
                    self.assertEqual(response.status_code, 403)
            response = self.client.get(
                f"/api/reports/boss/review-snapshots/{self.snapshot.share_code}/export/"
            )
            self.assertEqual(response.status_code, 403)
            workbook.assert_not_called()
        self.assertEqual(self.client.get("/api/reports/boss/receivables/").status_code, 200)
        self.assertEqual(
            self.client.get(
                f"/api/reports/boss/review-snapshots/{self.snapshot.share_code}/"
            ).status_code,
            200,
        )

    def test_exports_keep_valid_workbooks_and_warehouse_boundaries(self):
        response = self.client.post(
            "/api/reports/boss/exports/", {"report_type": "receivables"}, format="json"
        )
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), read_only=True)
        self.assertTrue(book.sheetnames)
        book.close()
        self.assertEqual(
            self.client.get(
                f"/api/reports/boss/review-snapshots/{self.snapshot.share_code}/export/"
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/reports/boss/exports/",
                {
                    "report_type": "receivables",
                    "warehouse": self.other_warehouse.pk,
                },
                format="json",
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.get(
                f"/api/reports/boss/review-snapshots/{self.other_snapshot.share_code}/export/"
            ).status_code,
            403,
        )

    def test_export_permission_alone_does_not_grant_boss_access(self):
        self.revoke("view_boss_dashboard")
        self.assertEqual(
            self.client.post(
                "/api/reports/boss/exports/", {"report_type": "receivables"}, format="json"
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.get(
                f"/api/reports/boss/review-snapshots/{self.snapshot.share_code}/export/"
            ).status_code,
            403,
        )

    def test_superuser_can_export_without_boss_role_or_group(self):
        user = get_user_model().objects.create_superuser(
            username="boss-export-root", email="qa@example.invalid", password="test"
        )
        self.client.force_authenticate(user)
        self.assertEqual(
            self.client.post(
                "/api/reports/boss/exports/", {"report_type": "receivables"}, format="json"
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.get(
                f"/api/reports/boss/review-snapshots/{self.other_snapshot.share_code}/export/"
            ).status_code,
            200,
        )
